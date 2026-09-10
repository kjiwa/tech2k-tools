from __future__ import annotations

import logging
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
from marcone.client import MarconeClient
from marcone.exceptions import MarconeError, PartNotFoundError

from inventory_updater.cache import PriceCache

logger = logging.getLogger(__name__)


@dataclass
class UpdateStats:
    total_rows: int = 0
    skipped: int = 0
    updated: int = 0
    not_found: int = 0
    errors: int = 0
    cancelled: bool = False


@dataclass
class RowUpdateResult:
    row_idx: int
    part_no: str
    status: str  # "updated", "not_found", "skipped", "error"
    old_cost: float | None = None
    new_cost: float | None = None
    old_price: float | None = None
    new_price: float | None = None
    message: str = ""


def get_file_info(file_path: str | Path) -> dict[str, Any]:
    """Inspect an Excel inventory file without modifying it."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = wb.active
    if sheet is None:
        raise ValueError(f"Workbook {path} has no active sheet")

    first_row = next(sheet.iter_rows(values_only=True), None)
    headers = list(first_row) if first_row else []
    col_map = InventoryUpdater._detect_columns(headers)

    row_count = sheet.max_row - 1 if sheet.max_row and sheet.max_row > 1 else 0

    wb.close()
    return {
        "file_name": path.name,
        "sheet_name": sheet.title,
        "row_count": max(0, row_count),
        "headers": [str(h) for h in headers if h is not None],
        "has_part_col": "part" in col_map,
        "has_cost_col": "cost" in col_map or "avg_cost" in col_map,
        "has_price_col": "price" in col_map,
        "has_supplier_col": "supplier" in col_map,
    }


class InventoryUpdater:
    """Updates part pricing in inventory spreadsheets using Marcone."""

    def __init__(
        self,
        client: MarconeClient,
        cache: PriceCache | None = None,
        cache_ttl_seconds: float = 86400 * 7,
    ) -> None:
        self.client = client
        self.cache = cache or PriceCache()
        self.cache_ttl = cache_ttl_seconds

    def update_file(
        self,
        input_file: str | Path,
        output_file: str | Path | None = None,
        field: str = "both",
        dry_run: bool = False,
        limit: int | None = None,
        supplier_filter: str | None = "Marcone",
        allow_blank_supplier: bool = True,
        set_supplier: str | None = None,
        only_missing: bool = False,
        progress_cb: Callable[[int, int, str], None] | None = None,
        row_cb: Callable[[RowUpdateResult], None] | None = None,
        cancel_check: Callable[[], bool] | None = None,
    ) -> UpdateStats:
        """Process an Excel inventory file and update prices."""
        input_path = Path(input_file)
        target_path = Path(output_file) if output_file else input_path

        wb = openpyxl.load_workbook(input_path)
        sheet = wb.active
        if sheet is None:
            raise ValueError(f"Workbook {input_file} has no active sheet")

        headers = [cell.value for cell in sheet[1]]
        col_map = self._detect_columns(headers)

        stats = UpdateStats()
        rows_to_process = list(range(2, sheet.max_row + 1))
        stats.total_rows = len(rows_to_process)

        processed_count = 0
        for row_idx in rows_to_process:
            if cancel_check and cancel_check():
                stats.cancelled = True
                break

            if limit is not None and stats.updated >= limit:
                break

            part_cell = sheet.cell(row=row_idx, column=col_map["part"])
            part_no = str(part_cell.value or "").strip()
            if not part_no:
                stats.skipped += 1
                if row_cb:
                    row_cb(
                        RowUpdateResult(
                            row_idx=row_idx,
                            part_no="",
                            status="skipped",
                            message="Blank part number",
                        )
                    )
                continue

            supplier_cell = (
                sheet.cell(row=row_idx, column=col_map["supplier"])
                if "supplier" in col_map
                else None
            )
            current_supplier = (
                str(supplier_cell.value or "").strip() if supplier_cell else ""
            )

            if supplier_filter:
                is_blank = not current_supplier
                is_match = supplier_filter.lower() in current_supplier.lower()
                if not (is_match or (allow_blank_supplier and is_blank)):
                    stats.skipped += 1
                    if row_cb:
                        row_cb(
                            RowUpdateResult(
                                row_idx=row_idx,
                                part_no=part_no,
                                status="skipped",
                                message=f"Supplier mismatch: '{current_supplier}'",
                            )
                        )
                    continue

            cost_cell = (
                sheet.cell(row=row_idx, column=col_map["cost"])
                if "cost" in col_map
                else None
            )
            avg_cost_cell = (
                sheet.cell(row=row_idx, column=col_map["avg_cost"])
                if "avg_cost" in col_map
                else None
            )
            price_cell = (
                sheet.cell(row=row_idx, column=col_map["price"])
                if "price" in col_map
                else None
            )

            def _parse_float(val: Any) -> float | None:
                try:
                    return float(val) if val is not None else None
                except (ValueError, TypeError):
                    return None

            orig_cost = _parse_float(
                cost_cell.value
                if cost_cell
                else (avg_cost_cell.value if avg_cost_cell else None)
            )
            orig_price = _parse_float(price_cell.value if price_cell else None)

            if only_missing:
                has_cost = any(
                    c is not None and c.value not in (None, "", 0, 0.0, "0", "0.00")
                    for c in (cost_cell, avg_cost_cell)
                )
                has_price = price_cell is not None and price_cell.value not in (
                    None,
                    "",
                    0,
                    0.0,
                    "0",
                    "0.00",
                )

                if field == "cost" and has_cost:
                    stats.skipped += 1
                    if row_cb:
                        row_cb(
                            RowUpdateResult(
                                row_idx=row_idx,
                                part_no=part_no,
                                status="skipped",
                                old_cost=orig_cost,
                                old_price=orig_price,
                                message="Cost already present",
                            )
                        )
                    continue
                if field == "price" and has_price:
                    stats.skipped += 1
                    if row_cb:
                        row_cb(
                            RowUpdateResult(
                                row_idx=row_idx,
                                part_no=part_no,
                                status="skipped",
                                old_cost=orig_cost,
                                old_price=orig_price,
                                message="Price already present",
                            )
                        )
                    continue
                if field == "both" and has_cost and has_price:
                    stats.skipped += 1
                    if row_cb:
                        row_cb(
                            RowUpdateResult(
                                row_idx=row_idx,
                                part_no=part_no,
                                status="skipped",
                                old_cost=orig_cost,
                                old_price=orig_price,
                                message="Cost and price already present",
                            )
                        )
                    continue

            processed_count += 1
            if progress_cb:
                progress_cb(processed_count, stats.total_rows, part_no)

            try:
                pricing = self._resolve_pricing(part_no)
                if not pricing or not pricing.has_pricing:
                    logger.warning("Part %s not found in Marcone catalog", part_no)
                    stats.not_found += 1
                    if row_cb:
                        row_cb(
                            RowUpdateResult(
                                row_idx=row_idx,
                                part_no=part_no,
                                status="not_found",
                                old_cost=orig_cost,
                                old_price=orig_price,
                                message="Not found in Marcone catalog",
                            )
                        )
                    continue

                modified = False
                new_cost = orig_cost
                new_price = orig_price

                if field in ("cost", "both") and pricing.customer_cost is not None:
                    if cost_cell and cost_cell.value != pricing.customer_cost:
                        cost_cell.value = pricing.customer_cost
                        modified = True
                        new_cost = pricing.customer_cost
                    if avg_cost_cell and avg_cost_cell.value != pricing.customer_cost:
                        avg_cost_cell.value = pricing.customer_cost
                        modified = True
                        new_cost = pricing.customer_cost

                if (
                    field in ("price", "both")
                    and price_cell
                    and pricing.list_price is not None
                    and price_cell.value != pricing.list_price
                ):
                    price_cell.value = pricing.list_price
                    modified = True
                    new_price = pricing.list_price

                if set_supplier and supplier_cell and modified:
                    supplier_cell.value = set_supplier

                if modified:
                    stats.updated += 1
                    if row_cb:
                        row_cb(
                            RowUpdateResult(
                                row_idx=row_idx,
                                part_no=part_no,
                                status="updated",
                                old_cost=orig_cost,
                                new_cost=new_cost,
                                old_price=orig_price,
                                new_price=new_price,
                            )
                        )
                else:
                    stats.skipped += 1
                    if row_cb:
                        row_cb(
                            RowUpdateResult(
                                row_idx=row_idx,
                                part_no=part_no,
                                status="skipped",
                                old_cost=orig_cost,
                                new_cost=orig_cost,
                                old_price=orig_price,
                                new_price=orig_price,
                                message="Price unchanged",
                            )
                        )

            except PartNotFoundError:
                logger.warning("Part %s not found in Marcone catalog", part_no)
                stats.not_found += 1
                if row_cb:
                    row_cb(
                        RowUpdateResult(
                            row_idx=row_idx,
                            part_no=part_no,
                            status="not_found",
                            old_cost=orig_cost,
                            old_price=orig_price,
                            message="Not found in Marcone catalog",
                        )
                    )
            except MarconeError as exc:
                logger.error("Error looking up part %s: %s", part_no, exc)
                stats.errors += 1
                if row_cb:
                    row_cb(
                        RowUpdateResult(
                            row_idx=row_idx,
                            part_no=part_no,
                            status="error",
                            old_cost=orig_cost,
                            old_price=orig_price,
                            message=str(exc),
                        )
                    )
            except Exception as exc:  # noqa: BLE001
                logger.error("Unexpected error looking up part %s: %s", part_no, exc)
                stats.errors += 1
                if row_cb:
                    row_cb(
                        RowUpdateResult(
                            row_idx=row_idx,
                            part_no=part_no,
                            status="error",
                            old_cost=orig_cost,
                            old_price=orig_price,
                            message=str(exc),
                        )
                    )

        if not dry_run and stats.updated > 0:
            target_path.parent.mkdir(parents=True, exist_ok=True)
            wb.save(target_path)
            logger.info("Saved updated workbook to %s", target_path)

        return stats

    def _resolve_pricing(self, part_no: str):
        """Check cache first, then query client and cache result."""
        cached = self.cache.get(part_no, max_age_seconds=self.cache_ttl)
        if cached:
            return cached

        if self.cache.is_known_missing(part_no, max_age_seconds=self.cache_ttl):
            return None

        try:
            pricing = self.client.lookup_part(part_no)
            self.cache.set(pricing)
            return pricing
        except PartNotFoundError:
            self.cache.set_not_found(part_no)
            return None

    @staticmethod
    def _detect_columns(headers: list[object]) -> dict[str, int]:
        """Map header names to 1-based column indices."""
        mapping: dict[str, int] = {}
        for idx, h in enumerate(headers, start=1):
            if not h:
                continue
            h_str = str(h).strip().lower()
            if (
                h_str in ("part no", "part no.", "part #", "partnumber", "part number")
                or h_str == "part"
            ):
                mapping.setdefault("part", idx)
            elif "vendor part" in h_str:
                mapping.setdefault("vendor_part", idx)
            elif h_str in ("avg. unit cost", "avg unit cost", "average unit cost"):
                mapping.setdefault("avg_cost", idx)
            elif h_str in ("purchase price", "supplier cost") or h_str == "cost":
                mapping.setdefault("cost", idx)
            elif h_str in ("unit price", "price *", "price"):
                mapping.setdefault("price", idx)
            elif h_str in ("primary vendor", "supplier name", "supplier", "vendor"):
                mapping.setdefault("supplier", idx)
            elif "manufacturer" in h_str or h_str == "make":
                mapping.setdefault("make", idx)

        mapping.setdefault("part", 1)
        if "cost" not in mapping and "avg_cost" in mapping:
            mapping["cost"] = mapping["avg_cost"]
        mapping.setdefault("cost", 8)
        mapping.setdefault("price", 10)
        mapping.setdefault("supplier", 14)
        return mapping
