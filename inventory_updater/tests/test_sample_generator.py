from pathlib import Path
from unittest.mock import MagicMock

import openpyxl
import pytest
from inventory_updater.updater import InventoryUpdater, MissingColumnsError, get_file_info

from scripts.generate_sample_xlsx import HEADERS, SAMPLE_ROWS, generate_sample_inventory

ROOT = Path(__file__).resolve().parent.parent.parent


def test_generate_sample_inventory(tmp_path):
    out_file = tmp_path / "test_sample.xlsx"
    path = generate_sample_inventory(out_file)
    assert path.exists()

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    assert ws.title == "Worksheet"
    assert ws.max_row == len(SAMPLE_ROWS) + 1
    assert ws.max_column == len(HEADERS)

    info = get_file_info(path)
    assert info["row_count"] == len(SAMPLE_ROWS)
    assert info["has_part_col"] is True
    assert info["has_cost_col"] is True
    assert info["has_price_col"] is True
    assert info["has_supplier_col"] is True
    assert "Marcone" in info["suppliers"]
    assert "Other Supplier" in info["suppliers"]


def test_committed_sample_inventory_exists():
    sample_file = ROOT / "samples" / "sample_inventory.xlsx"
    assert sample_file.exists(), f"Expected committed sample at {sample_file}"

    info = get_file_info(sample_file)
    assert info["row_count"] == len(SAMPLE_ROWS)
    assert info["has_part_col"] is True
    assert info["has_cost_col"] is True
    assert info["has_price_col"] is True


def test_generated_sample_has_all_required_columns(tmp_path):
    path = generate_sample_inventory(tmp_path / "sample.xlsx")
    assert get_file_info(path)["missing_columns"] == []


def test_sample_layout_without_header_row_is_refused(tmp_path):
    path = tmp_path / "headerless.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in SAMPLE_ROWS:
        ws.append(row)
    wb.save(path)
    before = path.read_bytes()
    updater = InventoryUpdater(client=MagicMock(), cache=MagicMock())

    with pytest.raises(MissingColumnsError):
        updater.update_file(input_file=path)

    assert path.read_bytes() == before
