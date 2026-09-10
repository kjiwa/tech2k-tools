from unittest.mock import MagicMock

import openpyxl
import pytest
from inventory_updater.cache import PriceCache
from inventory_updater.updater import InventoryUpdater
from marcone.client import MarconeClient
from marcone.exceptions import PartNotFoundError
from marcone.models import PartPricing


@pytest.fixture
def sample_excel(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Items"

    headers = [
        "Part No",
        "Item Name *",
        "Item Description",
        "Primary Category",
        "Secondary Category",
        "Barcode",
        "Vendor Part No",
        "Supplier Cost",
        "Cost Type",
        "Price *",
        "Price Type",
        "Taxable",
        "Manufacturer",
        "Supplier Name",
        "Status",
    ]
    ws.append(headers)
    ws.append(
        [
            "WPW10321304",
            "FILTER",
            "",
            "",
            "",
            "",
            "",
            10.0,
            "Flat",
            18.0,
            "Flat",
            "Yes",
            "WPL",
            "Marcone",
            "Active",
        ]
    )
    ws.append(
        [
            "240323002",
            "BIN",
            "",
            "",
            "",
            "",
            "",
            None,
            "Flat",
            None,
            "Flat",
            "Yes",
            "FRG",
            "Other",
            "Active",
        ]
    )
    ws.append(
        [
            "NOTFOUND99",
            "UNKNOWN",
            "",
            "",
            "",
            "",
            "",
            None,
            "Flat",
            None,
            "Flat",
            "Yes",
            "",
            "",
            "Active",
        ]
    )

    file_path = tmp_path / "InventoryItems.xlsx"
    wb.save(file_path)
    return file_path


def test_updater_both_fields(tmp_path, sample_excel):
    mock_client = MagicMock(spec=MarconeClient)

    def fake_lookup(part_no):
        if part_no == "WPW10321304":
            return PartPricing(
                part_number=part_no, customer_cost=15.25, list_price=27.50
            )
        elif part_no == "240323002":
            return PartPricing(
                part_number=part_no, customer_cost=18.45, list_price=30.00
            )
        raise PartNotFoundError(part_no)

    mock_client.lookup_part.side_effect = fake_lookup

    cache = PriceCache(db_path=str(tmp_path / "cache.sqlite"))
    updater = InventoryUpdater(client=mock_client, cache=cache)

    output_path = tmp_path / "InventoryItems_updated.xlsx"
    stats = updater.update_file(
        input_file=sample_excel, output_file=output_path, field="both"
    )

    assert stats.total_rows == 3
    assert stats.updated == 2
    assert stats.not_found == 1

    updated_wb = openpyxl.load_workbook(output_path)
    sheet = updated_wb.active

    # Row 2 (WPW10321304): cost col 8, price col 10
    assert sheet.cell(row=2, column=8).value == 15.25
    assert sheet.cell(row=2, column=10).value == 27.50

    # Row 3 (240323002)
    assert sheet.cell(row=3, column=8).value == 18.45
    assert sheet.cell(row=3, column=10).value == 30.00


def test_updater_cost_only(tmp_path, sample_excel):
    mock_client = MagicMock(spec=MarconeClient)
    mock_client.lookup_part.return_value = PartPricing(
        part_number="WPW10321304", customer_cost=15.25, list_price=27.50
    )

    cache = PriceCache(db_path=str(tmp_path / "cache.sqlite"))
    updater = InventoryUpdater(client=mock_client, cache=cache)

    output_path = tmp_path / "out.xlsx"
    updater.update_file(
        input_file=sample_excel, output_file=output_path, field="cost", limit=1
    )

    updated_wb = openpyxl.load_workbook(output_path)
    sheet = updated_wb.active
    assert sheet.cell(row=2, column=8).value == 15.25
    assert sheet.cell(row=2, column=10).value == 18.0  # unchanged list price


def test_updater_supplier_filter(tmp_path, sample_excel):
    mock_client = MagicMock(spec=MarconeClient)
    mock_client.lookup_part.return_value = PartPricing(
        part_number="WPW10321304", customer_cost=15.25, list_price=27.50
    )

    cache = PriceCache(db_path=str(tmp_path / "cache.sqlite"))
    updater = InventoryUpdater(client=mock_client, cache=cache)

    output_path = tmp_path / "out.xlsx"
    stats = updater.update_file(
        input_file=sample_excel,
        output_file=output_path,
        supplier_filter="Marcone",
    )

    assert stats.updated == 1
    assert stats.skipped == 2


def test_updater_set_supplier(tmp_path, sample_excel):
    mock_client = MagicMock(spec=MarconeClient)
    mock_client.lookup_part.return_value = PartPricing(
        part_number="240323002", customer_cost=18.45, list_price=30.00
    )

    cache = PriceCache(db_path=str(tmp_path / "cache.sqlite"))
    updater = InventoryUpdater(client=mock_client, cache=cache)

    output_path = tmp_path / "out.xlsx"
    updater.update_file(
        input_file=sample_excel,
        output_file=output_path,
        only_missing=True,
        set_supplier="Marcone Supply",
    )

    updated_wb = openpyxl.load_workbook(output_path)
    sheet = updated_wb.active
    # Row 3 supplier name is at column 14
    assert sheet.cell(row=3, column=14).value == "Marcone Supply"


def test_updater_dry_run(tmp_path, sample_excel):
    mock_client = MagicMock(spec=MarconeClient)
    mock_client.lookup_part.return_value = PartPricing(
        part_number="WPW10321304", customer_cost=15.25, list_price=27.50
    )

    cache = PriceCache(db_path=str(tmp_path / "cache.sqlite"))
    updater = InventoryUpdater(client=mock_client, cache=cache)

    output_path = tmp_path / "should_not_exist.xlsx"
    stats = updater.update_file(
        input_file=sample_excel,
        output_file=output_path,
        dry_run=True,
        limit=1,
    )

    assert stats.updated == 1
    assert not output_path.exists()


@pytest.fixture
def service_fusion_excel(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Worksheet"

    headers = [
        "Product Category",
        "Product Name",
        "Part Number",
        "Model",
        "SKU",
        "Type",
        "Quantity On Hand",
        "Unit Price",
        "Avg. Unit Cost",
        "Member Price",
        "Active?",
        "Inventory Item?",
        "QB Class",
        "Media URL",
        "Pays Commission",
        "Fixed Commission",
        "Percentage Comission",
        "Pays Bonus",
        "Fixed Bonus",
        "Percentage Bonus",
        "Pays Hours",
        "Primary Vendor",
        "Purchase Price",
        "Secondary Vendor",
        "Purchase Price",
        "Tertiary Vendor",
        "Purchase Price",
        "Related Products",
        "Purchase Description",
        "Description",
    ]
    ws.append(headers)
    ws.append(
        [
            "Parts",
            "FILTER",
            "WPW10321304",
            "",
            "",
            "",
            "1",
            18.0,
            10.0,
            "",
            "Yes",
            "Yes",
            "",
            "",
            "No",
            0,
            0,
            "No",
            0,
            0,
            0,
            "Marcone",
            10.0,
            "",
            0,
            "",
            0,
            "",
            "",
            "FILTER",
        ]
    )
    ws.append(
        [
            "Parts",
            "BIN",
            "240323002",
            "",
            "",
            "",
            "0",
            0.0,
            0.0,
            "",
            "Yes",
            "Yes",
            "",
            "",
            "No",
            0,
            0,
            "No",
            0,
            0,
            0,
            "",
            0.0,
            "",
            0,
            "",
            0,
            "",
            "",
            "BIN",
        ]
    )

    file_path = tmp_path / "Service_Fusion_Inventory.xlsx"
    wb.save(file_path)
    return file_path


def test_updater_service_fusion_format(tmp_path, service_fusion_excel):
    mock_client = MagicMock(spec=MarconeClient)

    def fake_lookup(part_no):
        if part_no == "WPW10321304":
            return PartPricing(
                part_number=part_no, customer_cost=15.25, list_price=27.50
            )
        elif part_no == "240323002":
            return PartPricing(
                part_number=part_no, customer_cost=18.45, list_price=30.00
            )
        raise PartNotFoundError(part_no)

    mock_client.lookup_part.side_effect = fake_lookup

    cache = PriceCache(db_path=str(tmp_path / "cache.sqlite"))
    updater = InventoryUpdater(client=mock_client, cache=cache)

    output_path = tmp_path / "Service_Fusion_Inventory_updated.xlsx"
    stats = updater.update_file(
        input_file=service_fusion_excel,
        output_file=output_path,
        field="both",
        set_supplier="Marcone",
    )

    assert stats.total_rows == 2
    assert stats.updated == 2

    updated_wb = openpyxl.load_workbook(output_path)
    sheet = updated_wb.active

    # Row 2 (WPW10321304): unit_price col 8, avg_cost col 9, purchase_price col 23, vendor col 22
    assert sheet.cell(row=2, column=8).value == 27.50
    assert sheet.cell(row=2, column=9).value == 15.25
    assert sheet.cell(row=2, column=23).value == 15.25
    assert sheet.cell(row=2, column=22).value == "Marcone"

    # Row 3 (240323002)
    assert sheet.cell(row=3, column=8).value == 30.00
    assert sheet.cell(row=3, column=9).value == 18.45
    assert sheet.cell(row=3, column=23).value == 18.45
    assert sheet.cell(row=3, column=22).value == "Marcone"
